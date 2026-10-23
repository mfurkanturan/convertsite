"""Independent content/structure checks. Unmeasured gates are explicit, never a pass."""
import hashlib
import io
import json
import re
import subprocess
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from difflib import SequenceMatcher
import pdfplumber
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageStat


def command(args):
    return subprocess.run(list(map(str,args)), capture_output=True, timeout=40, check=True).stdout


def normalized(text):
    return ''.join(unicodedata.normalize('NFC', text).split())


def edit_distance(a, b):
    if a==b: return 0
    if not a: return len(b)
    if not b: return len(a)
    previous = list(range(len(b)+1))
    for i, ca in enumerate(a, 1):
        current=[i]
        for j, cb in enumerate(b, 1): current.append(min(current[-1]+1,previous[j]+1,previous[j-1]+(ca!=cb)))
        previous=current
    return previous[-1]


def accuracy(expected, actual):
    # Levenshtein, bounded synthetic documents. CER, not SequenceMatcher similarity.
    a,b=normalized(expected),normalized(actual)
    return max(0,1-edit_distance(a,b)/max(1,len(a)))


def pdf_text(path):
    return command(['/usr/bin/pdftotext', '-layout', path, '-']).decode('utf-8')


def reference_text(path):
    if path.name!='layout.pdf': return pdf_text(path)
    # Declared logical column reading order from layout-fixture.py, not geometric row interleaving.
    lines=['Deniz Örnek','Yazılım Mühendisi','example.com']
    for section in range(1,6):
        lines.append(f'Deneyim {section}')
        lines.extend(f'Çözüm {section}.{line}: İstanbul için güvenilir belge işleme sistemleri geliştirdi.' for line in range(1,6))
    for label in ['Sol','Sağ']:
        lines.append(f'{label} Sütun')
        lines.extend(f'{label} bölüm satır {line}: Türkçe içerik.' for line in range(1,16))
    return '\n'.join(lines)


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        root=ET.fromstring(z.read('word/document.xml'))
        return ' '.join(n.text or '' for n in root.iter() if n.tag.endswith('}t'))


def pdf_summary(path):
    with pdfplumber.open(path) as pdf:
        return [{'width':p.width,'height':p.height,'rotation':p.rotation,'text':normalized(command(['/usr/bin/pdftotext','-f',i,'-l',i,path,'-']).decode('utf-8'))} for i,p in enumerate(pdf.pages,1)]


def compare_visual(source, target, directory):
    results=[]
    for label, path in [('source',source),('output',target)]:
        command(['/usr/bin/pdftoppm','-r','72','-png',path,directory/label])
    for original in sorted(directory.glob('source-*.png')):
        converted=directory/original.name.replace('source-','output-')
        if not converted.exists(): results.append({'page':original.name,'same_size':False}); continue
        with Image.open(original) as a, Image.open(converted) as b:
            if a.size!=b.size: results.append({'page':original.name,'same_size':False}); continue
            diff=ImageChops.difference(a.convert('RGB'),b.convert('RGB'))
            results.append({'page':original.name,'same_size':True,'mean_pixel_error':sum(ImageStat.Stat(diff).mean)/3/255})
            diff.save(directory/original.name.replace('source-','diff-'))
    return results


def evaluate(case, output, kind, corpus, directory):
    checks={}; metrics={}; missing=[]
    tool=case['tool']; source=corpus/case['inputs'][0]; expect=case['expect']
    text=''; signature={}
    if kind=='docx':
        text=docx_text(output); d=Document(output)
        checks['editable_text']=bool(normalized(text))
        metrics['content_accuracy']=accuracy(reference_text(source),text)
        checks['content_accuracy']=metrics['content_accuracy']>=.995
        checks['critical_numbers']=re.findall(r'\d+(?:[.,]\d+)*',normalized(reference_text(source)))==re.findall(r'\d+(?:[.,]\d+)*',normalized(text))
        converted=directory/'word-render'; converted.mkdir()
        command(['/usr/bin/libreoffice',f'-env:UserInstallation={(directory/"profile").as_uri()}','--headless','--convert-to','pdf','--outdir',converted,output])
        rendered=converted/(output.stem+'.pdf')
        original=pdf_summary(source); actual=pdf_summary(rendered)
        checks['page_count']=len(original)==len(actual)
        metrics['visual']=compare_visual(source,rendered,directory)
        displacements=[]
        with pdfplumber.open(source) as a, pdfplumber.open(rendered) as b:
            for pa,pb in zip(a.pages,b.pages):
                aw,bw=pa.extract_words(),pb.extract_words()
                for word in aw:
                    matches=[w for w in bw if w['text']==word['text']]
                    if matches: displacements.append(min(max(abs(w['x0']-word['x0']),abs(w['top']-word['top'])) for w in matches))
        metrics['max_word_displacement_pt']=max(displacements,default=None)
        checks['layout_5pt']=bool(displacements) and max(displacements)<=5
        missing += ['manual_clipping_overlap_review','Microsoft_Word_rendering']
        signature={'text':normalized(text),'pages':actual}
    elif kind=='xlsx':
        w=load_workbook(output); cells=[[str(c.value or '') for row in sheet for c in row] for sheet in w]
        signature={'cells':cells}; checks['editable_cells']=bool(cells)
        checks['no_executable_formulas']=all(c.data_type!='f' for sheet in w for row in sheet for c in row)
        # Ruled table fixture has independently declared ground truth.
        expected=[['Ürün','Adet','Not'],['Kalem','12','=1+1'],['Defter','7','İstanbul']]
        if source.name in ['table.pdf','borderless.pdf','merged.pdf']:
            if source.name!='table.pdf': expected=[['Kalem','12','1234,56'] for _ in range(4)]
            observed=[[str(c.value or '') for c in row] for row in w.worksheets[0]]
            exact=sum(i<len(observed) and j<len(observed[i]) and normalized(v)==normalized(observed[i][j]) for i,row in enumerate(expected) for j,v in enumerate(row))
            count=sum(map(len,expected))
            metrics['cell_accuracy']=exact/count; checks['cell_accuracy']=exact/count>=.99
            checks['structure']=len(observed)==len(expected) and all(len(row)==3 for row in observed)
            checks['critical_numbers']=checks['structure'] and all(observed[i][j]==value for i,row in enumerate(expected) for j,value in enumerate(row) if any(c.isdigit() for c in value))
        elif source.name=='multi-table.pdf':
            expected_sheets=[]
            for page in range(2): expected_sheets.append([['Tarih','Tutar','Kod'],[f'{23+page*2}.10.2026','1234,56' if page==0 else '0,00',f'00{12+page*2}'],[f'{24+page*2}.10.2026','-45,10' if page==0 else '9999,99',f'00{13+page*2}']])
            observed=[[[str(c.value or '') for c in row] for row in sheet] for sheet in w]
            exact=sum(i<len(observed) and r<len(observed[i]) and col<len(observed[i][r]) and value==observed[i][r][col] for i,sheet in enumerate(expected_sheets) for r,row in enumerate(sheet) for col,value in enumerate(row))
            metrics['cell_accuracy']=exact/18; checks['cell_accuracy']=exact/18>=.99
            checks['structure']=len(observed)==2 and all(len(sheet)==3 and all(len(row)==3 for row in sheet) for sheet in observed)
            checks['critical_numbers']=exact==18
        else: missing += ['independent_cell_ground_truth','merged_cell_structure']
        metrics['numeric_type_inference']=False  # Current contract deliberately writes extracted values as inert strings.
    elif kind=='txt':
        text=output.read_text(); expected='HELLO WORLD London England DOCUMENT TEST 2026' if source.name=='english.jpg' else 'MERHABA DUNYA İstanbul Türkiye BELGE TESTI 2026'
        metrics['character_accuracy']=accuracy(expected,text)
        expected_words=expected.split(); observed_words=text.split()
        metrics['word_error_rate']=edit_distance(expected_words,observed_words)/max(1,len(expected_words))
        metrics['word_sequence_similarity']=SequenceMatcher(None,expected_words,observed_words).ratio()
        checks['ocr_accuracy']=metrics['character_accuracy']>=(.95 if source.stem in ['blur','skew','small','rotate'] else .98)
        signature={'text':normalized(text)}
    elif kind=='pdf':
        command(['/usr/bin/qpdf','--check',output]); text=pdf_text(output)
        summary=pdf_summary(output); signature={'pages':summary}
        if expect.get('image_pages'): checks['page_count']=len(summary)==expect['image_pages']; missing += ['image_orientation_color_crop_review']
        elif expect.get('office'):
            ext=expect['office']
            if ext=='docx': expected=docx_text(source)
            elif ext=='xlsx': expected='Kalem'
            else: expected='İstanbul Türkiye Şişli öğle çözüm 1234,56 2026-10-23'
            checks['source_text_present']=normalized(expected) in normalized(text)
            if ext=='pptx': checks['slide_count']=len(summary)==(2 if int(source.stem.split('-')[-1])%2 else 1)
            if ext=='xlsx': checks['calculated_value']=bool(re.search(r'\b5\b',text)); missing+=['print_layout_reference']
            missing+=['manual_clipping_overlap_review','source_application_reference']
        elif expect.get('numbers'):
            checks['all_page_numbers']=all(normalized(f'{i} / {len(summary)}') in p['text'] for i,p in enumerate(summary,1))
            missing+=['number_position_geometry']
        elif expect.get('overlay'):
            checks['all_watermarks']=all(normalized(expect['overlay']) in p['text'] for p in summary)
            missing+=['watermark_position_geometry']
        elif expect.get('selection'):
            plan=[]
            for part in expect['selection'].split(','):
                bounds=list(map(int,part.split('-'))); first,last=bounds[0],bounds[-1]
                plan+=list(range(first,last+(1 if last>=first else -1),1 if last>=first else -1))
            checks['order']=len(summary)==len(plan) and all(f'SAYFA{i}' in p['text'] for i,p in zip(plan,summary))
            checks['rotation']=all(p['rotation']==expect['rotation'] for p in summary)
        elif tool=='merge-pdf':
            expected=''.join(normalized(pdf_text(corpus/f)) for f in case['inputs'])
            checks['content_order']=normalized(text)==expected
            checks['pages']=len(summary)==sum(len(pdf_summary(corpus/f)) for f in case['inputs'])
        elif tool in ['compress-pdf','unlock-pdf']:
            original=corpus/expect.get('source',source.name)
            checks['content_preserved']=normalized(text)==normalized(pdf_text(original))
            metrics['visual']=compare_visual(original,output,directory)
            checks['lossless_visual']=all(p.get('mean_pixel_error')==0 for p in metrics['visual'])
            if tool=='compress-pdf': checks['not_larger']=output.stat().st_size<=source.stat().st_size; metrics['size_ratio']=output.stat().st_size/source.stat().st_size
    elif kind in ['png','jpg','zip']:
        images=[]
        if kind=='zip':
            with zipfile.ZipFile(output) as z:
                for name in z.namelist():
                    data=z.read(name)
                    if name.endswith('.pdf'):
                        member=directory/('verify-'+Path(name).name); member.write_bytes(data)
                        images.append({'name':name,'pages':pdf_summary(member)})
                    else:
                        with Image.open(io.BytesIO(data)) as im: images.append({'name':name,'size':im.size,'format':im.format})
        else:
            with Image.open(output) as im: images.append({'size':im.size,'format':im.format})
        signature={'members':images}
        if tool=='pdf-to-images':
            checks['page_count']=len(images)==len(pdf_summary(source))
            checks['dimensions']=all(max(i['size'])<=3000 for i in images)
            checks['format']=all(i['format']==('JPEG' if case['options']['format']=='jpg' else 'PNG') for i in images)
            missing+=['raster_content_color_review']
        else:
            selected=list(map(int,expect['selection'].split(',')))
            checks['split_count']=len(images)==len(selected)
            checks['split_order']=len(images)==len(selected) and all(len(item['pages'])==1 and f'SAYFA{n}' in item['pages'][0]['text'] for n,item in zip(selected,images))
            checks['split_rotation']=all(item['pages'][0]['rotation']==expect['rotation'] for item in images)
    if not checks: missing.append('no_acceptance_checks')
    return dict(checks=checks,metrics=metrics,unmeasured=missing,signature=hashlib.sha256(json.dumps(signature,sort_keys=True).encode()).hexdigest(),passed=bool(checks) and all(checks.values()))
