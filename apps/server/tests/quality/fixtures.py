"""Synthetic corpus only; never reads user documents."""
import json
import runpy
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from openpyxl import Workbook

FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
TEXT = 'İstanbul Türkiye Şişli öğle çözüm 1234,56 2026-10-23'
PPTX_GENERATOR = None


def pptx(path, text, count):
    if PPTX_GENERATOR is None:
        raise RuntimeError('Quality PPTX fixtures require the validated engine fixture generator.')
    PPTX_GENERATOR(path, [f'{text} SLIDE {i}' for i in range(1, count + 1)])


def generate(root):
    root.mkdir(parents=True, exist_ok=True)
    previous = sys.argv
    try:
        sys.argv = ['engine-fixtures.py', str(root)]
        engine_fixtures = runpy.run_path(str(Path(__file__).parent.parent / 'engine-fixtures.py'), run_name='__main__')
        global PPTX_GENERATOR
        PPTX_GENERATOR = engine_fixtures['make_pptx']
    finally:
        sys.argv = previous
    pdfmetrics.registerFont(TTFont('Audit', FONT))
    for index, name in enumerate(['plain', 'unicode', 'borderless', 'rotated', 'multipage', 'landscape', 'blank', 'merged']):
        c = canvas.Canvas(str(root / f'{name}.pdf'), pagesize=(842, 595) if name == 'landscape' else (595, 842))
        for page in range(3 if name == 'multipage' else 1):
            c.setFont('Audit', 12)
            if name == 'rotated':
                c.setPageRotation(90)
            if name != 'blank':
                c.drawString(40, 500, TEXT if name != 'plain' else 'REPORT 1234.56 TEST 2026')
                c.drawString(40, 470, f'PAGE {page + 1}')
            if name in ('borderless', 'merged'):
                if name == 'merged':
                    # Four real table rows; the first row spans the first two columns.
                    for y in [310, 340, 370, 400]: c.line(35, y, 490, y)
                    for x in [35, 185, 335, 490]: c.line(x, 310, x, 400)
                    c.line(35, 400, 35, 430); c.line(490, 400, 490, 430)
                    c.line(335, 400, 335, 430)
                    c.line(35, 430, 490, 430)
                    c.drawString(45, 415, 'Kalem')
                    c.drawString(345, 415, '1234,56')
                    for row, y in enumerate([385, 355, 325], 1):
                        for col, value in enumerate(['Kalem', '12', '1234,56']):
                            c.drawString(45 + col * 150, y, value)
                else:
                    for row in range(4):
                        for col, value in enumerate(['Kalem', '12', '1234,56']):
                            c.drawString(40 + col * 150, 400 - row * 30, value)
            c.showPage()
        c.save()
        d = Document(); d.add_heading(f'Audit {index}', 0); d.add_paragraph(TEXT)
        if index % 2: d.add_table(rows=2, cols=3).cell(0, 0).text = 'Kalem 1234,56'
        if index == 4: d.add_page_break(); d.add_paragraph('PAGE 2')
        d.save(root / f'office-{index}.docx')
        w = Workbook(); w.active.append(['Ürün', 'Adet', 'Tutar']); w.active.append(['Kalem', 12, 1234.56])
        w.active.append(['Formül', '=2+3', '=SUM(B2:B3)'])
        if index % 2: w.create_sheet('İkinci').append(['SECOND', 42])
        if index == 5:
            for r in range(100): w.active.append(['Satır', r, r * 1.5])
        w.save(root / f'office-{index}.xlsx')
        pptx(root / f'office-{index}.pptx', TEXT, 2 if index % 2 else 1)
    base = Image.open(root / 'text.png')
    exif=base.getexif(); exif[274]=6
    base.save(root/'exif.jpg',quality=95,exif=exif)
    english=Image.new('RGB',(1800,650),'white'); draw=ImageDraw.Draw(english)
    font=ImageFont.truetype(FONT,64)
    for index,line in enumerate(['HELLO WORLD','London England','DOCUMENT TEST 2026']): draw.text((80,60+index*150),line,fill='black',font=font)
    english.save(root/'english.jpg',quality=95)
    c=canvas.Canvas(str(root/'multi-table.pdf'))
    for page in range(2):
        xs=[40,210,360,500]; ys=[700,670,640,610]
        for x in xs: c.line(x,610,x,700)
        for y in ys: c.line(40,y,500,y)
        rows=[['Tarih','Tutar','Kod'],[f'{23+page*2}.10.2026','1234,56' if page==0 else '0,00',f'00{12+page*2}'],[f'{24+page*2}.10.2026','-45,10' if page==0 else '9999,99',f'00{13+page*2}']]
        c.setFont('Audit',11)
        for r,row in enumerate(rows):
            for col,value in enumerate(row): c.drawString(xs[col]+5,ys[r]-17,value)
        c.showPage()
    c.save()
    for name, im in [('blur', base.filter(ImageFilter.GaussianBlur(1))), ('skew', base.rotate(3, expand=True, fillcolor='white')), ('rotate', base.rotate(90, expand=True)), ('small', base.resize((900, 325))), ('transparent', base.convert('RGBA'))]:
        if name=='transparent': ImageDraw.Draw(im).rectangle((0,0,250,40),fill=(255,0,0,0))
        im.save(root / f'{name}.png')
    (root / 'broken.pdf').write_bytes(b'%PDF-1.7\nbroken')
    c = canvas.Canvas(str(root/'ocr-long.pdf'), pagesize=(900,325))
    for i in range(19):
        c.drawImage(str(root/'text.png'),0,0,900,325); c.showPage()
    c.save()
    # Resource tiers: bounded pages, no personal information, comfortably below 10 MiB.
    for size, count in [('small', 1), ('medium', 10), ('near', 299)]:
        c = canvas.Canvas(str(root / f'perf-{size}.pdf'))
        for i in range(count):
            c.setFont('Audit', 12); c.drawString(40, 700, TEXT); c.drawString(40, 670, f'PAGE {i+1}')
            c.showPage()
        c.save()
    return root


def performance_input(tool, tier, corpus, base):
    counts={'small':1,'medium':10,'near':49 if tool in ['pdf-to-word','pdf-to-excel'] else 19 if tool in ['ocr','images-to-pdf'] else 29 if tool=='pdf-to-images' else 296 if tool=='merge-pdf' else 299}
    count=counts[tier]; options=dict(base['options'])
    if tool=='images-to-pdf': return ['text.png']*count,options,count
    if tool in ['word-to-pdf','excel-to-pdf','pptx-to-pdf']:
        ext={'word-to-pdf':'docx','excel-to-pdf':'xlsx','pptx-to-pdf':'pptx'}[tool]
        path=corpus/f'perf-{tool}-{tier}.{ext}'
        if ext=='docx':
            d=Document()
            for i in range(count):
                if i: d.add_page_break()
                d.add_paragraph(TEXT)
            d.save(path)
        elif ext=='xlsx':
            w=Workbook(); w.remove(w.active)
            for i in range(count): w.create_sheet(f'Sheet{i+1}').append([TEXT,1234.56,'=2+3'])
            w.save(path); options={'sheets':'all'}
        else: pptx(path,TEXT,count)
        return [path.name],options,count
    path=corpus/f'perf-{tool}-{tier}.pdf'
    if tool=='pdf-to-excel':
        c=canvas.Canvas(str(path))
        for i in range(count):
            for x in [40,180,300]: c.line(x,600,x,700)
            for y in [600,650,700]: c.line(40,y,300,y)
            c.setFont('Audit',12); c.drawString(50,675,'Kalem'); c.drawString(190,675,'1234,56'); c.drawString(50,625,'Defter'); c.drawString(190,625,'12'); c.showPage()
        c.save()
    elif tool=='ocr':
        c=canvas.Canvas(str(path),pagesize=(900,325))
        for i in range(count): c.drawImage(str(corpus/'text.png'),0,0,900,325); c.showPage()
        c.save()
    else:
        import subprocess
        subprocess.run(['/usr/bin/qpdf',str(corpus/'perf-near.pdf'),'--pages','.',f'1-{count}','--',str(path)],check=True)
    if tool=='unlock-pdf':
        import subprocess
        encrypted=corpus/f'encrypted-perf-{tier}.pdf'
        subprocess.run(['/usr/bin/qpdf',str(path),'--encrypt','AuditSecret123','AuditSecret123','256','--',str(encrypted)],check=True); path=encrypted
    if tool=='edit-pdf-pages': options={'pages':f'1-{count}','mode':'pdf','rotation':90}
    return [path.name]+(['pages.pdf'] if tool=='merge-pdf' else []),options,count
