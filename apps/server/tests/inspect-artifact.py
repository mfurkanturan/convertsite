import io
import json
import sys
import zipfile
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
import pdfplumber

kind, path = sys.argv[1:]
result = {}
if kind == "docx":
    document = Document(path)
    paragraphs = [p for p in document.paragraphs] + [p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    result = {
        "text": "\n".join(p.text for p in paragraphs),
        "runs": [{"text": r.text, "bold": r.bold, "italic": r.italic, "size": r.font.size.pt if r.font.size else None} for p in paragraphs for r in p.runs],
        "images": len(list(document.element.iter(qn('w:drawing')))),
        "externalLinks": sum(rel.is_external for part in document.part.package.parts for rel in part.rels.values()),
        "sectionWidths": [section.page_width.pt for section in document.sections],
    }
elif kind == "pdf":
    with pdfplumber.open(path) as pdf:
        result = {"pages": [{"width": page.width, "height": page.height, "words": page.extract_words(), "images": [{"x0": image['x0'], "top": image['top']} for image in page.images]} for page in pdf.pages]}
elif kind == "xlsx":
    workbook = load_workbook(path, data_only=False)
    result = {"sheets": workbook.sheetnames, "cells": [{"value": cell.value, "type": cell.data_type} for sheet in workbook for row in sheet for cell in row]}
elif kind == "zip":
    members = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            data = archive.read(name)
            if name.endswith(".pdf"):
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    members.append({"name": name, "pages": len(pdf.pages), "width": pdf.pages[0].width})
            else:
                with Image.open(io.BytesIO(data)) as image:
                    members.append({"name": name, "width": image.width, "height": image.height, "format": image.format})
    result = {"members": members}
elif kind in ("png", "jpg"):
    with Image.open(path) as image:
        result = {"format": image.format, "width": image.width, "height": image.height}
print(json.dumps(result, ensure_ascii=False))
