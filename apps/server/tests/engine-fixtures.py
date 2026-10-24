import sys
import runpy
import zipfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

directory = Path(sys.argv[1])
directory.mkdir(exist_ok=True, parents=True)
runpy.run_path(str(Path(__file__).with_name("layout-fixture.py")))["create_layout_fixture"](directory)
font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
pdfmetrics.registerFont(TTFont("FixtureFont", font_path))
pdf = canvas.Canvas(str(directory / "table.pdf"))
pdf.setFont("FixtureFont", 16)
pdf.drawString(40, 760, "Türkçe Rapor: İstanbul")
xs, ys = [40, 180, 280, 430], [700, 675, 650, 625]
for x in xs:
    pdf.line(x, ys[-1], x, ys[0])
for y in ys:
    pdf.line(xs[0], y, xs[-1], y)
rows = [["Ürün", "Adet", "Not"], ["Kalem", "12", "=1+1"], ["Defter", "7", "İstanbul"]]
pdf.setFont("FixtureFont", 11)
for r, row in enumerate(rows):
    for c, value in enumerate(row):
        pdf.drawString(xs[c] + 5, ys[r] - 17, value)
pdf.save()
pdf = canvas.Canvas(str(directory / "pages.pdf"))
for index, width in enumerate([300, 400, 500], 1):
    pdf.setPageSize((width, 500))
    pdf.setFont("FixtureFont", 18)
    pdf.drawString(30, 450, f"SAYFA {index}")
    pdf.showPage()
pdf.save()
image = Image.new("RGB", (1800, 650), "white")
draw = ImageDraw.Draw(image)
font = ImageFont.truetype(font_path, 64)
for index, text in enumerate(["MERHABA DUNYA", "İstanbul Türkiye", "BELGE TESTI 2026"]):
    draw.text((80, 60 + index * 150), text, fill="black", font=font)
image.save(directory / "text.png")
image.save(directory / "text.jpg", quality=95)
pdf = canvas.Canvas(str(directory / "scanned.pdf"), pagesize=(900, 325))
pdf.drawImage(str(directory / "text.png"), 0, 0, 900, 325)
pdf.save()
document = Document()
document.add_heading("Türkçe Belge", 0)
document.add_paragraph("Merhaba İstanbul. Şifre çözümü değil, dönüşüm testi.")
document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Kalem"
document.save(directory / "document.docx")
workbook = Workbook()
first = workbook.active
first.title = "Birinci"
first.append(["Ürün", "Adet"])
first.append(["İstanbul Kalem", 12])
second = workbook.create_sheet("İkinci")
second.append(["İkinci Sayfa", 2026])
for sheet in workbook:
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
workbook.save(directory / "workbook.xlsx")
with zipfile.ZipFile(directory / "document.docx") as source:
    for bad, addition in [("macro.docx", "word/vbaProject.bin"), ("traversal.docx", "../escape.xml")]:
        with zipfile.ZipFile(directory / bad, "w", zipfile.ZIP_DEFLATED) as target:
            for name in source.namelist():
                target.writestr(name, source.read(name))
            target.writestr(addition, "unsafe fixture")
workbook.active["A1"].hyperlink = "https://example.com"
workbook.save(directory / "external.xlsx")
Image.new("RGB", (12000, 1)).save(directory / "wide.png")


def make_pptx(path, slides, hidden=(), extra=None):
    """Valid synthetic PPTX (master, layout, theme) so LibreOffice Impress opens it; slides is a list of texts."""
    from xml.sax.saxutils import escape
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pk = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct = "application/vnd.openxmlformats-officedocument.presentationml"
    count = len(slides)
    tree = ('<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
            '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>')
    fills = "".join('<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>' for _ in range(3))
    lines = "".join('<a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln>' for _ in range(3))
    theme = (f'<a:theme xmlns:a="{a}" name="Fixture"><a:themeElements><a:clrScheme name="c">'
             + "".join(f'<a:{n}><a:srgbClr val="{v}"/></a:{n}>' for n, v in [("dk1", "000000"), ("lt1", "FFFFFF"), ("dk2", "1F2937"), ("lt2", "F3F4F6"), ("accent1", "2563EB"), ("accent2", "DC2626"), ("accent3", "16A34A"), ("accent4", "D97706"), ("accent5", "7C3AED"), ("accent6", "0891B2"), ("hlink", "2563EB"), ("folHlink", "7C3AED")])
             + '</a:clrScheme><a:fontScheme name="f"><a:majorFont><a:latin typeface="DejaVu Sans"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>'
             '<a:minorFont><a:latin typeface="DejaVu Sans"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont></a:fontScheme>'
             f'<a:fmtScheme name="s"><a:fillStyleLst>{fills}</a:fillStyleLst><a:lnStyleLst>{lines}</a:lnStyleLst>'
             f'<a:effectStyleLst>{"<a:effectStyle><a:effectLst/></a:effectStyle>" * 3}</a:effectStyleLst><a:bgFillStyleLst>{fills}</a:bgFillStyleLst></a:fmtScheme></a:themeElements></a:theme>')
    overrides = [("/ppt/presentation.xml", f"{ct}.presentation.main+xml"), ("/ppt/slideMasters/slideMaster1.xml", f"{ct}.slideMaster+xml"),
                 ("/ppt/slideLayouts/slideLayout1.xml", f"{ct}.slideLayout+xml"), ("/ppt/theme/theme1.xml", "application/vnd.openxmlformats-officedocument.theme+xml")]
    overrides += [(f"/ppt/slides/slide{i}.xml", f"{ct}.slide+xml") for i in range(1, count + 1)]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        extras = extra or {}
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                   + "".join(f'<Override PartName="{n}" ContentType="{t}"/>' for n, t in overrides) + "</Types>")
        z.writestr("_rels/.rels", f'<Relationships xmlns="{pk}"><Relationship Id="rId1" Type="{r}/officeDocument" Target="ppt/presentation.xml"/></Relationships>')
        ids = "".join(f'<p:sldId id="{255 + i}" r:id="rId{i + 1}"/>' for i in range(1, count + 1))
        z.writestr("ppt/presentation.xml", f'<p:presentation xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}"><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst><p:sldIdLst>{ids}</p:sldIdLst><p:sldSz cx="9144000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>')
        z.writestr("ppt/_rels/presentation.xml.rels", f'<Relationships xmlns="{pk}"><Relationship Id="rId1" Type="{r}/slideMaster" Target="slideMasters/slideMaster1.xml"/>'
                   + "".join(f'<Relationship Id="rId{i + 1}" Type="{r}/slide" Target="slides/slide{i}.xml"/>' for i in range(1, count + 1)) + f'<Relationship Id="rId{count + 2}" Type="{r}/theme" Target="theme/theme1.xml"/></Relationships>')
        z.writestr("ppt/theme/theme1.xml", theme)
        z.writestr("ppt/slideMasters/slideMaster1.xml", f'<p:sldMaster xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}"><p:cSld><p:spTree>{tree}</p:spTree></p:cSld><p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/><p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst></p:sldMaster>')
        z.writestr("ppt/slideMasters/_rels/slideMaster1.xml.rels", f'<Relationships xmlns="{pk}"><Relationship Id="rId1" Type="{r}/slideLayout" Target="../slideLayouts/slideLayout1.xml"/><Relationship Id="rId2" Type="{r}/theme" Target="../theme/theme1.xml"/></Relationships>')
        z.writestr("ppt/slideLayouts/slideLayout1.xml", f'<p:sldLayout xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}" type="blank"><p:cSld name="Blank"><p:spTree>{tree}</p:spTree></p:cSld></p:sldLayout>')
        z.writestr("ppt/slideLayouts/_rels/slideLayout1.xml.rels", f'<Relationships xmlns="{pk}"><Relationship Id="rId1" Type="{r}/slideMaster" Target="../slideMasters/slideMaster1.xml"/></Relationships>')
        for i, text in enumerate(slides, 1):
            show = ' show="0"' if i in hidden else ""
            z.writestr(f"ppt/slides/slide{i}.xml", f'<p:sld xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}"{show}><p:cSld><p:spTree>{tree}<p:sp><p:nvSpPr><p:cNvPr id="2" name="Metin"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="500000" y="500000"/><a:ext cx="8000000" cy="1200000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr wrap="square"/><a:lstStyle/><a:p><a:r><a:rPr lang="tr-TR" sz="2400"/><a:t>{escape(text)}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
            relationship = f"ppt/slides/_rels/slide{i}.xml.rels"
            default = f'<Relationships xmlns="{pk}"><Relationship Id="rId1" Type="{r}/slideLayout" Target="../slideLayouts/slideLayout1.xml"/></Relationships>'
            z.writestr(relationship, extras.pop(relationship, default))
        for name, content in extras.items():
            z.writestr(name, content)


make_pptx(directory / "deck.pptx", ["Türkçe Sunum: İstanbul Şişli", "İkinci slayt: öğle çözüm ğüşiöç", "Üçüncü slayt GİZLİ"], hidden=(3,))
make_pptx(directory / "macro.pptx", ["Makro"], extra={"ppt/vbaProject.bin": "unsafe fixture"})
make_pptx(directory / "traversal.pptx", ["Yol"], extra={"../escape.xml": "unsafe fixture"})
make_pptx(directory / "external.pptx", ["Dış"], extra={"ppt/slides/_rels/slide1.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com" TargetMode="External"/></Relationships>'})
with zipfile.ZipFile(directory / "notpptx.pptx", "w") as z:
    z.writestr("hello.txt", "zip but not a presentation")


import fitz
for name, rotation, crop in [("wm-rotated90.pdf", 90, None), ("wm-rotated270.pdf", 270, None), ("wm-cropped.pdf", 0, (100, 100, 400, 500))]:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 120), "Kaynak metin korunmalı", fontsize=14)
    if crop:
        page.set_cropbox(fitz.Rect(*crop))
    page.set_rotation(rotation)
    doc.save(directory / name)
    doc.close()


for name, rotation, vertical in [("word-rot90.pdf", 90, False), ("word-rot180.pdf", 180, False), ("word-rot270.pdf", 270, False), ("word-content90.pdf", 0, True)]:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    for index, line in enumerate(["DONUK SAYFA 1234,56", "İstanbul Türkiye 2026", "Üçüncü satır öğle çözüm"]):
        if vertical:
            page.insert_text((100 + index * 30, 700), line, fontname="Audit", fontfile=font_path, fontsize=12, rotate=90)
        else:
            page.insert_text((50, 100 + index * 30), line, fontname="Audit", fontfile=font_path, fontsize=12)
    page.set_rotation(rotation)
    doc.save(directory / name)
    doc.close()


# Multi-row paragraphs whose source pitch (30 pt, 18 pt) differs from the converter's default, with a first line
# that fills the text width so a trailing space plus line break would add an empty line.
pdf = canvas.Canvas(str(directory / "pitch.pdf"), pagesize=(595, 842))
for page, (pitch, size) in enumerate([(30, 12), (18, 11)], 1):
    pdf.setFont("FixtureFont", size)
    pdf.drawString(40, 760, "Satır adımı sayfa %d: İstanbul Türkiye Şişli öğle çözüm 1234,56 2026-10-23 uzun satır" % page)
    for row in range(1, 5):
        pdf.drawString(40, 760 - row * pitch, "Satır %d öğütücü çözüm %d" % (row, page * 100 + row))
    pdf.showPage()
pdf.save()
