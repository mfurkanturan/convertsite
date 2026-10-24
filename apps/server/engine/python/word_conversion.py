"""Reconstruct editable Word layout from PDF geometry, rather than plain text."""
import logging
import pymupdf
from docx import Document
from docx.oxml.ns import qn
from common import EngineError, input_path
from word_spacing import install as install_line_spacing
from word_layout import explicit_column_breaks, positioned_images, table_grids, trim_line_break_spaces


def remove_external_links(document):
    # Keep styled link text, without creating active external Office relationships.
    # Uploaded Office documents are subject to the same external-link policy.
    for part in document.part.package.parts:
        external = {key for key, relation in part.rels.items() if relation.is_external}
        if not external:
            continue
        root = getattr(part, "element", None)
        if root is None:
            raise EngineError("INVALID_DOCUMENT")
        for link in list(root.iter(qn("w:hyperlink"))):
            if link.get(qn("r:id")) not in external:
                continue
            parent = link.getparent()
            position = parent.index(link)
            for child in list(link):
                # pdf2docx nests hyperlinks inside an already styled w:r.
                # A w:r inside another w:r is invalid and disappears in Writer.
                content = list(child) if parent.tag == qn("w:r") and child.tag == qn("w:r") else [child]
                for element in content:
                    if parent.tag == qn("w:r") and element.tag == qn("w:rPr"):
                        continue
                    parent.insert(position, element)
                    position += 1
            parent.remove(link)
        for key in external:
            part.drop_rel(key)


def upright_rotation(page):
    # pdf2docx only keeps horizontal text. Pick the page rotation under which the dominant text reads left to right.
    weight = {}
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            key = (round(line["dir"][0]), round(line["dir"][1]))
            weight[key] = weight.get(key, 0) + sum(len(span["text"]) for span in line["spans"])
    if not weight:
        return page.rotation
    dx, dy = max(weight, key=weight.get)
    for rotation in (0, 90, 180, 270):
        matrix = pymupdf.Matrix(1, 0, 0, 1, 0, 0).prerotate(rotation)
        if round(dx * matrix.a + dy * matrix.c) == 1:
            return rotation
    return page.rotation


def normalized_copy(source, directory):
    # Working copy only; the uploaded file is never modified. Returns the path to convert from.
    with pymupdf.open(source) as pdf:
        changed = False
        for page in pdf:
            target = upright_rotation(page)
            if target != page.rotation:
                page.set_rotation(target)
                changed = True
        if not changed:
            return source
        path = directory / "upright.pdf"
        pdf.save(path)
        return path


def page_text(source):
    with pymupdf.open(source) as pdf:
        return "".join(page.get_text() for page in pdf)


def alphanumeric(text):
    return sum(character.isalnum() for character in text)


def docx_characters(document):
    paragraphs = list(document.paragraphs) + [p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    return alphanumeric("".join(p.text for p in paragraphs))


def pdf_to_word(inputs, options, directory):
    # Lazy import keeps the layout engine's native dependencies out of other tools.
    from pdf2docx import Converter
    install_line_spacing()

    source = input_path(inputs[0])
    with pymupdf.open(source) as pdf:
        characters = sum(len(page.get_text()) for page in pdf)
    if characters > 1_000_000:
        raise EngineError("LIMIT_EXCEEDED")
    if not characters:
        raise EngineError("NO_TEXT")

    expected = alphanumeric(page_text(source))
    source = normalized_copy(source, directory)
    output = directory / "result.docx"
    # The converter uses the root logger for progress; never log document content.
    logging.disable(logging.CRITICAL)
    converter = Converter(str(source))
    try:
        # Keep physical line breaks so short column lines cannot merge and overflow.
        converter.convert(str(output), multi_processing=False, ignore_page_error=False,
                          line_break_width_ratio=1.0)
    finally:
        converter.close()
    document = Document(output)
    # A converted document that lost most of the source text is a failure, not a successful empty file.
    if docx_characters(document) < expected * 0.8:
        raise EngineError("INVALID_DOCUMENT")
    remove_external_links(document)
    explicit_column_breaks(document)
    table_grids(document)
    trim_line_break_spaces(document)
    with pymupdf.open(source) as pdf:
        positioned_images(document, pdf)
    document.save(output)
    return output, "docx"
