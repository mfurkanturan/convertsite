import re
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from common import input_path, EngineError
from excel_tables import page_tables, merged_ranges


def clean_text(value):
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    if len(cleaned) > 32767:
        raise EngineError("LIMIT_EXCEEDED")
    return cleaned


def pdf_to_excel(inputs, options, directory):
    workbook = Workbook()
    workbook.remove(workbook.active)
    count = 0
    total_cells = 0
    with pdfplumber.open(input_path(inputs[0])) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            for rows, source_table in page_tables(page):
                count += 1
                sheet = workbook.create_sheet(f"Sayfa{page_number}_Tablo{count}")
                for row_number, row in enumerate(rows, 1):
                    total_cells += len(row)
                    if total_cells > 100000 or len(row) > 100 or row_number > 20000:
                        raise EngineError("LIMIT_EXCEEDED")
                    for column, value in enumerate(row, 1):
                        cell = sheet.cell(row_number, column, clean_text(value))
                        # Extracted formulas are untrusted text, never executable formulas.
                        cell.data_type = "s"
                        if row_number == 1:
                            cell.font = Font(bold=True, color="FFFFFF")
                            cell.fill = PatternFill("solid", fgColor="3564F4")
                for row_number, first_column, last_column in merged_ranges(source_table):
                    sheet.merge_cells(start_row=row_number, start_column=first_column,
                                      end_row=row_number, end_column=last_column)
                sheet.freeze_panes = "A2"
                for column in range(1, sheet.max_column + 1):
                    sheet.column_dimensions[get_column_letter(column)].width = 24
    if not count:
        raise EngineError("NO_TABLES")
    output = directory / "result.xlsx"
    workbook.save(output)
    return output, "xlsx"
